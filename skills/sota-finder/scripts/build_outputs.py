#!/usr/bin/env python3
"""Normalize a SOTA research bundle and build the four standard artifacts."""

from __future__ import annotations

import argparse
import copy
import html
import json
import re
import shutil
import sys
import tempfile
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

SCHEMA_VERSION = "sota-finder.evidence.v1"
DEFAULT_ARTIFACTS = (
    "brief-report.md",
    "full-report.html",
    "leaderboards.xlsx",
    "evidence.json",
)
AUTHORITATIVE_SOURCE_ROLES = {"official", "primary", "author"}
SOURCE_ROLES = AUTHORITATIVE_SOURCE_ROLES | {"secondary", "aggregator"}
SECTIONS = {"main", "reference", "unavailable"}
CONFLICT_STATUSES = {"resolved", "unresolved", "unavailable"}


class ContractError(ValueError):
    """Raised when the prepared research input violates a mechanical invariant."""


def _now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise ContractError(message)


def _text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return str(value).strip()


def slugify(value: Any, fallback: str = "item", max_length: int = 60) -> str:
    """Return a stable filesystem/ID fragment without random or cryptographic data."""
    value = unicodedata.normalize("NFKD", _text(value)).encode("ascii", "ignore").decode("ascii").lower()
    chars: list[str] = []
    previous_dash = False
    for char in value:
        if char.isalnum():
            chars.append(char)
            previous_dash = False
        elif not previous_dash:
            chars.append("-")
            previous_dash = True
    result = "".join(chars).strip("-") or fallback
    result = re.sub(r"-+", "-", result)
    return result[:max_length].rstrip("-") or fallback


def _unique_id(prefix: str, seed: Any, used: set[str]) -> str:
    base = f"{prefix}-{slugify(seed)}"
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}-{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def _resolve(ref: Any, mapping: dict[str, str], kind: str) -> str:
    key = _text(ref)
    _require(key in mapping, f"Unknown {kind} reference: {key!r}")
    return mapping[key]


def _resolve_many(values: Any, mapping: dict[str, str], kind: str) -> list[str]:
    if values is None:
        return []
    _require(isinstance(values, list), f"{kind} references must be a list")
    resolved: list[str] = []
    for value in values:
        item = _resolve(value, mapping, kind)
        if item not in resolved:
            resolved.append(item)
    return resolved


def _normalize_metric(scope: dict[str, Any]) -> None:
    metric = scope.get("metric")
    _require(isinstance(metric, dict), f"Scope {scope.get('name')!r} requires a metric object")
    metric["name"] = _text(metric.get("name"))
    _require(metric["name"], f"Scope {scope.get('name')!r} requires metric.name")
    direction = _text(metric.get("direction")).lower()
    _require(direction in {"higher", "lower"}, f"Metric direction must be 'higher' or 'lower', got {direction!r}")
    metric["direction"] = direction
    metric.setdefault("variant", None)
    metric.setdefault("unit", None)


def _normalize_score(claim: dict[str, Any]) -> None:
    numeric = claim.get("score_numeric")
    if numeric is not None:
        _require(isinstance(numeric, (int, float)) and not isinstance(numeric, bool), f"Claim {claim.get('method')!r} score_numeric must be a number or null")
        claim["score_numeric"] = float(numeric)
    display = _text(claim.get("score_display"))
    if not display and numeric is not None:
        display = f"{numeric:g}"
    claim["score_display"] = display or "not available"


def normalize_bundle(raw: dict[str, Any]) -> dict[str, Any]:
    """Assign stable IDs, resolve local keys, rank main claims, and return final evidence data."""
    _require(isinstance(raw, dict), "Research input must be a JSON object")
    data = copy.deepcopy(raw)

    run = data.get("research_run")
    _require(isinstance(run, dict), "research_run must be an object")
    topic = _text(run.get("topic"))
    query = _text(run.get("query"))
    cutoff = _text(run.get("evidence_checked_through"))
    language = _text(run.get("language"), "en")
    _require(topic, "research_run.topic is required")
    _require(query, "research_run.query is required")
    _require(re.fullmatch(r"\d{4}-\d{2}-\d{2}", cutoff) is not None, "research_run.evidence_checked_through must be YYYY-MM-DD")
    try:
        datetime.strptime(cutoff, "%Y-%m-%d")
    except ValueError as exc:
        raise ContractError("research_run.evidence_checked_through is not a valid calendar date") from exc
    topic_slug = slugify(run.get("topic_slug") or topic, "sota-leaderboard")
    run["topic"] = topic
    run["topic_slug"] = topic_slug
    run["query"] = query
    run["evidence_checked_through"] = cutoff
    run["generated_at"] = _text(run.get("generated_at")) or _now_iso()
    run["language"] = language
    run["run_id"] = _text(run.get("run_id")) or f"run-{topic_slug}-{cutoff.replace('-', '')}"
    aminer = run.setdefault("aminer", {"used": False})
    _require(isinstance(aminer, dict) and isinstance(aminer.get("used"), bool), "research_run.aminer.used must be boolean")

    coverage = data.setdefault("coverage", {})
    _require(isinstance(coverage, dict), "coverage must be an object")
    coverage.setdefault("searched_source_types", [])
    coverage.setdefault("queries", [])
    coverage.setdefault("stop_reason", "")
    coverage.setdefault("gaps", [])
    for field in ("searched_source_types", "queries", "gaps"):
        _require(isinstance(coverage[field], list), f"coverage.{field} must be a list")
    _require(_text(coverage.get("stop_reason")), "coverage.stop_reason is required")

    scopes = data.get("scopes")
    sources = data.get("sources")
    claims = data.get("claims")
    conflicts = data.setdefault("conflicts", [])
    _require(isinstance(scopes, list) and scopes, "scopes must be a non-empty list")
    _require(isinstance(sources, list), "sources must be a list")
    _require(isinstance(claims, list), "claims must be a list")
    _require(isinstance(conflicts, list), "conflicts must be a list")

    scope_ids: set[str] = set()
    scope_map: dict[str, str] = {}
    for index, scope in enumerate(scopes, 1):
        _require(isinstance(scope, dict), f"Scope #{index} must be an object")
        local_key = _text(scope.get("key") or scope.get("scope_id") or scope.get("name") or f"scope-{index}")
        name = _text(scope.get("name"))
        task = _text(scope.get("task"))
        _require(name and task, f"Scope {local_key!r} requires name and task")
        scope_id = _text(scope.get("scope_id")) or _unique_id("scope", local_key, scope_ids)
        _require(scope_id.startswith("scope-"), f"scope_id must start with 'scope-': {scope_id}")
        _require(scope_id not in scope_map.values(), f"Duplicate scope_id: {scope_id}")
        scope["scope_id"] = scope_id
        scope["name"] = name
        scope["task"] = task
        scope.setdefault("benchmark", None)
        scope.setdefault("benchmark_version", None)
        scope.setdefault("dataset_or_subtask", None)
        scope.setdefault("split", None)
        scope.setdefault("evaluation_protocol", None)
        scope.setdefault("material_axes", [])
        _require(isinstance(scope["material_axes"], list), f"Scope {name!r} material_axes must be a list")
        _normalize_metric(scope)
        for ref in (local_key, scope_id, name):
            if ref:
                _require(ref not in scope_map or scope_map[ref] == scope_id, f"Ambiguous scope reference: {ref!r}")
                scope_map[ref] = scope_id
        for field in ("main_claim_ids", "reference_claim_ids", "unavailable_claim_ids", "conflict_ids"):
            scope[field] = []
        scope.pop("key", None)

    source_ids: set[str] = set()
    source_map: dict[str, str] = {}
    for index, source in enumerate(sources, 1):
        _require(isinstance(source, dict), f"Source #{index} must be an object")
        local_key = _text(source.get("key") or source.get("source_id") or source.get("url") or source.get("title") or f"source-{index}")
        title = _text(source.get("title"))
        authority = _text(source.get("authority")).lower()
        accessed = _text(source.get("accessed_date"))
        _require(title, f"Source {local_key!r} requires title")
        _require(authority in SOURCE_ROLES, f"Source {local_key!r} has invalid authority: {authority!r}")
        _require(re.fullmatch(r"\d{4}-\d{2}-\d{2}", accessed) is not None, f"Source {local_key!r} accessed_date must be YYYY-MM-DD")
        _require(isinstance(source.get("available"), bool), f"Source {local_key!r} available must be boolean")
        source_id = _text(source.get("source_id")) or _unique_id("src", local_key, source_ids)
        _require(source_id.startswith("src-"), f"source_id must start with 'src-': {source_id}")
        _require(source_id not in source_map.values(), f"Duplicate source_id: {source_id}")
        source["source_id"] = source_id
        source["title"] = title
        source["authority"] = authority
        source["accessed_date"] = accessed
        source.setdefault("type", "other")
        source.setdefault("url", None)
        for ref in (local_key, source_id, _text(source.get("url")), title):
            if ref:
                _require(ref not in source_map or source_map[ref] == source_id, f"Ambiguous source reference: {ref!r}")
                source_map[ref] = source_id
        source.pop("key", None)

    claim_ids: set[str] = set()
    claim_map: dict[str, str] = {}
    for index, claim in enumerate(claims, 1):
        _require(isinstance(claim, dict), f"Claim #{index} must be an object")
        local_key = _text(claim.get("key") or claim.get("claim_id") or f"{claim.get('method', 'claim')}-{index}")
        method = _text(claim.get("method"))
        section = _text(claim.get("section")).lower()
        _require(method, f"Claim {local_key!r} requires method")
        _require(section in SECTIONS, f"Claim {local_key!r} has invalid section: {section!r}")
        scope_ref = claim.pop("scope_key", None) or claim.get("scope_id")
        scope_id = _resolve(scope_ref, scope_map, "scope")
        raw_source_refs = claim.pop("evidence_source_keys", None)
        if raw_source_refs is None:
            raw_source_refs = claim.get("evidence_source_ids", [])
        evidence_source_ids = _resolve_many(raw_source_refs, source_map, "source")
        claim_id = _text(claim.get("claim_id")) or _unique_id("clm", local_key, claim_ids)
        _require(claim_id.startswith("clm-"), f"claim_id must start with 'clm-': {claim_id}")
        _require(claim_id not in claim_map.values(), f"Duplicate claim_id: {claim_id}")
        claim["claim_id"] = claim_id
        claim["scope_id"] = scope_id
        claim["section"] = section
        claim["method"] = method
        claim.setdefault("variant", None)
        claim.setdefault("paper_title", None)
        claim.setdefault("year", None)
        claim.setdefault("setting", {})
        claim.setdefault("evidence_locators", [])
        claim.setdefault("method_summary", None)
        claim.setdefault("notes", None)
        _require(isinstance(claim.get("comparability_complete"), bool), f"Claim {local_key!r} comparability_complete must be boolean")
        _require(isinstance(claim["setting"], dict), f"Claim {local_key!r} setting must be an object")
        _require(isinstance(claim["evidence_locators"], list), f"Claim {local_key!r} evidence_locators must be a list")
        for locator in claim["evidence_locators"]:
            _require(isinstance(locator, dict), f"Claim {local_key!r} evidence locator must be an object")
            locator_ref = locator.pop("source_key", None) or locator.get("source_id")
            if locator_ref:
                locator["source_id"] = _resolve(locator_ref, source_map, "source")
        claim["evidence_source_ids"] = evidence_source_ids
        _normalize_score(claim)
        claim["rank"] = None
        claim["sota"] = False
        for ref in (local_key, claim_id):
            _require(ref not in claim_map or claim_map[ref] == claim_id, f"Ambiguous claim reference: {ref!r}")
            claim_map[ref] = claim_id
        claim.pop("key", None)

    conflict_ids: set[str] = set()
    conflict_map: dict[str, str] = {}
    for index, conflict in enumerate(conflicts, 1):
        _require(isinstance(conflict, dict), f"Conflict #{index} must be an object")
        local_key = _text(conflict.get("key") or conflict.get("conflict_id") or f"conflict-{index}")
        summary = _text(conflict.get("summary"))
        status = _text(conflict.get("status")).lower()
        _require(summary, f"Conflict {local_key!r} requires summary")
        _require(status in CONFLICT_STATUSES, f"Conflict {local_key!r} has invalid status: {status!r}")
        scope_ref = conflict.pop("scope_key", None) or conflict.get("scope_id")
        scope_id = _resolve(scope_ref, scope_map, "scope")
        claim_refs = conflict.pop("claim_keys", None)
        if claim_refs is None:
            claim_refs = conflict.get("claim_ids", [])
        source_refs = conflict.pop("source_keys", None)
        if source_refs is None:
            source_refs = conflict.get("source_ids", [])
        conflict_id = _text(conflict.get("conflict_id")) or _unique_id("conflict", local_key, conflict_ids)
        _require(conflict_id.startswith("conflict-"), f"conflict_id must start with 'conflict-': {conflict_id}")
        _require(conflict_id not in conflict_map.values(), f"Duplicate conflict_id: {conflict_id}")
        conflict["conflict_id"] = conflict_id
        conflict["scope_id"] = scope_id
        conflict["summary"] = summary
        conflict["status"] = status
        conflict.setdefault("resolution", None)
        conflict["claim_ids"] = _resolve_many(claim_refs, claim_map, "claim")
        conflict["source_ids"] = _resolve_many(source_refs, source_map, "source")
        conflict_map[local_key] = conflict_id
        conflict_map[conflict_id] = conflict_id
        conflict.pop("key", None)

    scopes_by_id = {scope["scope_id"]: scope for scope in scopes}
    claims_by_scope: dict[str, list[dict[str, Any]]] = {scope_id: [] for scope_id in scopes_by_id}
    for claim in claims:
        claims_by_scope[claim["scope_id"]].append(claim)

    ordered_claims: list[dict[str, Any]] = []
    for scope in scopes:
        scope_id = scope["scope_id"]
        direction = scope["metric"]["direction"]
        scoped = claims_by_scope[scope_id]
        main = [claim for claim in scoped if claim["section"] == "main"]
        reference = [claim for claim in scoped if claim["section"] == "reference"]
        unavailable = [claim for claim in scoped if claim["section"] == "unavailable"]
        for claim in main:
            _require(claim.get("score_numeric") is not None, f"Main claim {claim['claim_id']} requires score_numeric")
        main.sort(key=lambda claim: ((-claim["score_numeric"] if direction == "higher" else claim["score_numeric"]), claim["method"].casefold(), _text(claim.get("variant")).casefold()))
        previous_score: float | None = None
        previous_rank = 0
        for position, claim in enumerate(main, 1):
            score = claim["score_numeric"]
            if previous_score is None or score != previous_score:
                previous_rank = position
                previous_score = score
            claim["rank"] = previous_rank
            claim["sota"] = previous_rank == 1
        reference.sort(key=lambda claim: (claim["method"].casefold(), _text(claim.get("variant")).casefold()))
        unavailable.sort(key=lambda claim: (claim["method"].casefold(), _text(claim.get("variant")).casefold()))
        scope["main_claim_ids"] = [claim["claim_id"] for claim in main]
        scope["reference_claim_ids"] = [claim["claim_id"] for claim in reference]
        scope["unavailable_claim_ids"] = [claim["claim_id"] for claim in unavailable]
        scope["conflict_ids"] = [conflict["conflict_id"] for conflict in conflicts if conflict["scope_id"] == scope_id]
        ordered_claims.extend(main + reference + unavailable)

    data["schema_version"] = SCHEMA_VERSION
    data["scopes"] = scopes
    data["sources"] = sorted(sources, key=lambda source: source["source_id"])
    data["claims"] = ordered_claims
    data["conflicts"] = sorted(conflicts, key=lambda conflict: conflict["conflict_id"])
    warnings = validate_evidence(data)
    data["validation_warnings"] = warnings
    return data


def validate_evidence(data: dict[str, Any]) -> list[str]:
    """Check cross-file and main-gate invariants that do not require research judgment."""
    _require(data.get("schema_version") == SCHEMA_VERSION, f"schema_version must be {SCHEMA_VERSION}")
    scopes = data.get("scopes", [])
    claims = data.get("claims", [])
    sources = data.get("sources", [])
    conflicts = data.get("conflicts", [])
    _require(scopes, "At least one scope is required")

    scope_ids = [scope.get("scope_id") for scope in scopes]
    claim_ids = [claim.get("claim_id") for claim in claims]
    source_ids = [source.get("source_id") for source in sources]
    conflict_ids = [conflict.get("conflict_id") for conflict in conflicts]
    for name, ids in (("scope", scope_ids), ("claim", claim_ids), ("source", source_ids), ("conflict", conflict_ids)):
        _require(all(isinstance(value, str) and value for value in ids), f"Every {name} requires a non-empty ID")
        _require(len(ids) == len(set(ids)), f"Duplicate {name} IDs detected")

    scope_set = set(scope_ids)
    claim_set = set(claim_ids)
    source_set = set(source_ids)
    conflict_set = set(conflict_ids)
    source_by_id = {source["source_id"]: source for source in sources}
    claim_by_id = {claim["claim_id"]: claim for claim in claims}

    warnings: list[str] = []
    for scope in scopes:
        _require(scope["metric"]["direction"] in {"higher", "lower"}, f"Invalid metric direction for {scope['scope_id']}")
        membership = scope["main_claim_ids"] + scope["reference_claim_ids"] + scope["unavailable_claim_ids"]
        _require(all(claim_id in claim_set for claim_id in membership), f"Scope {scope['scope_id']} references an unknown claim")
        _require(all(conflict_id in conflict_set for conflict_id in scope["conflict_ids"]), f"Scope {scope['scope_id']} references an unknown conflict")
        for claim_id in membership:
            _require(claim_by_id[claim_id]["scope_id"] == scope["scope_id"], f"Claim {claim_id} is linked to the wrong scope")

    for claim in claims:
        _require(claim["scope_id"] in scope_set, f"Claim {claim['claim_id']} has unknown scope")
        _require(claim["section"] in SECTIONS, f"Claim {claim['claim_id']} has invalid section")
        _require(all(source_id in source_set for source_id in claim["evidence_source_ids"]), f"Claim {claim['claim_id']} references an unknown source")
        if claim["section"] == "main":
            _require(claim["comparability_complete"], f"Main claim {claim['claim_id']} must declare comparability_complete=true")
            _require(claim.get("score_numeric") is not None, f"Main claim {claim['claim_id']} requires score_numeric")
            _require(claim["evidence_source_ids"], f"Main claim {claim['claim_id']} requires evidence")
            authorities = {source_by_id[source_id]["authority"] for source_id in claim["evidence_source_ids"]}
            _require(authorities & AUTHORITATIVE_SOURCE_ROLES, f"Main claim {claim['claim_id']} requires official, primary, or author evidence")
            _require(isinstance(claim.get("rank"), int) and claim["rank"] >= 1, f"Main claim {claim['claim_id']} requires a positive rank")
        else:
            _require(claim.get("rank") is None and not claim.get("sota"), f"Non-main claim {claim['claim_id']} must not be ranked or marked SOTA")
        if not claim["evidence_source_ids"]:
            warnings.append(f"Claim {claim['claim_id']} has no linked evidence source.")

    for conflict in conflicts:
        _require(conflict["scope_id"] in scope_set, f"Conflict {conflict['conflict_id']} has unknown scope")
        _require(conflict["status"] in CONFLICT_STATUSES, f"Conflict {conflict['conflict_id']} has invalid status")
        _require(all(claim_id in claim_set for claim_id in conflict["claim_ids"]), f"Conflict {conflict['conflict_id']} references an unknown claim")
        _require(all(source_id in source_set for source_id in conflict["source_ids"]), f"Conflict {conflict['conflict_id']} references an unknown source")
        if conflict["status"] == "resolved" and not _text(conflict.get("resolution")):
            warnings.append(f"Resolved conflict {conflict['conflict_id']} has no resolution text.")

    return warnings


def _labels(language: str) -> dict[str, str]:
    if language.lower().startswith("zh"):
        return {
            "title": "SOTA 与 Leaderboard 报告",
            "cutoff": "证据核验截止",
            "main": "主排行榜",
            "reference": "参考结果",
            "unavailable": "冲突 / 不可用",
            "scope": "Scope",
            "metric": "指标",
            "method": "方法",
            "variant": "变体",
            "score": "分数",
            "rank": "排名",
            "evidence": "证据",
            "notes": "说明",
            "methods": "SOTA 方法简介",
            "conflicts": "冲突",
            "coverage": "覆盖与缺口",
            "sources": "证据来源",
            "gaps": "缺口",
            "stop": "停止依据",
            "filter": "筛选表格",
            "none": "无",
        }
    return {
        "title": "SOTA and Leaderboard Report",
        "cutoff": "Evidence checked through",
        "main": "Main Leaderboard",
        "reference": "Reference Results",
        "unavailable": "Conflicts / Unavailable",
        "scope": "Scope",
        "metric": "Metric",
        "method": "Method",
        "variant": "Variant",
        "score": "Score",
        "rank": "Rank",
        "evidence": "Evidence",
        "notes": "Notes",
        "methods": "SOTA Method Notes",
        "conflicts": "Conflicts",
        "coverage": "Coverage and Gaps",
        "sources": "Evidence Sources",
        "gaps": "Gaps",
        "stop": "Stop reason",
        "filter": "Filter tables",
        "none": "None",
    }


def _metric_text(scope: dict[str, Any]) -> str:
    metric = scope["metric"]
    parts = [metric["name"]]
    if metric.get("variant"):
        parts.append(_text(metric["variant"]))
    if metric.get("unit"):
        parts.append(f"[{metric['unit']}]")
    parts.append(f"({metric['direction']})")
    return " ".join(parts)


def _setting_text(setting: dict[str, Any]) -> str:
    return "; ".join(f"{key}={str(value).lower() if isinstance(value, bool) else value}" for key, value in sorted(setting.items()))


def _md_cell(value: Any) -> str:
    return _text(value, "—").replace("|", "\\|").replace("\n", " ") or "—"


def _claims_for_scope(data: dict[str, Any], scope_id: str, section: str) -> list[dict[str, Any]]:
    return [claim for claim in data["claims"] if claim["scope_id"] == scope_id and claim["section"] == section]


def _markdown_table(claims: list[dict[str, Any]], labels: dict[str, str], include_rank: bool) -> str:
    if not claims:
        return f"_{labels['none']}_\n"
    headers = ([labels["rank"]] if include_rank else []) + [labels["method"], labels["variant"], labels["score"], labels["evidence"], labels["notes"], "claim_id"]
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for claim in claims:
        cells: list[Any] = []
        if include_rank:
            cells.append(claim.get("rank"))
        cells.extend([
            claim["method"],
            claim.get("variant"),
            claim["score_display"],
            ", ".join(claim["evidence_source_ids"]),
            claim.get("notes"),
            claim["claim_id"],
        ])
        lines.append("| " + " | ".join(_md_cell(cell) for cell in cells) + " |")
    return "\n".join(lines) + "\n"


def render_brief_markdown(data: dict[str, Any]) -> str:
    run = data["research_run"]
    labels = _labels(run["language"])
    lines = [f"# {labels['title']}: {run['topic']}", "", f"**{labels['cutoff']}: {run['evidence_checked_through']}**", ""]
    for scope in data["scopes"]:
        lines.extend([f"## {scope['name']}", "", f"- {labels['metric']}: `{_metric_text(scope)}`"])
        if scope.get("evaluation_protocol"):
            lines.append(f"- Evaluation protocol: {_text(scope['evaluation_protocol'])}")
        if scope.get("material_axes"):
            axes = "; ".join(f"{_text(axis.get('name'))}={_text(axis.get('value'))}" for axis in scope["material_axes"])
            lines.append(f"- Material axes: {axes}")
        lines.extend(["", f"### {labels['main']}", "", _markdown_table(_claims_for_scope(data, scope["scope_id"], "main"), labels, True)])
        reference = _claims_for_scope(data, scope["scope_id"], "reference")
        if reference:
            lines.extend([f"### {labels['reference']}", "", _markdown_table(reference, labels, False)])
        unavailable = _claims_for_scope(data, scope["scope_id"], "unavailable")
        if unavailable:
            lines.extend([f"### {labels['unavailable']}", "", _markdown_table(unavailable, labels, False)])
        sota_claims = [claim for claim in _claims_for_scope(data, scope["scope_id"], "main") if claim.get("sota")]
        summaries = [claim for claim in sota_claims if _text(claim.get("method_summary"))]
        if summaries:
            lines.extend([f"### {labels['methods']}", ""])
            for claim in summaries:
                lines.append(f"- **{claim['method']}**: {_text(claim['method_summary'])} (`{claim['claim_id']}`)")
            lines.append("")
        scoped_conflicts = [conflict for conflict in data["conflicts"] if conflict["scope_id"] == scope["scope_id"]]
        if scoped_conflicts:
            lines.extend([f"### {labels['conflicts']}", ""])
            for conflict in scoped_conflicts:
                resolution = f" — {_text(conflict.get('resolution'))}" if _text(conflict.get("resolution")) else ""
                lines.append(f"- **{conflict['status']}**: {conflict['summary']}{resolution} (`{conflict['conflict_id']}`)")
            lines.append("")
    lines.extend([f"## {labels['coverage']}", "", f"- {labels['stop']}: {_text(data['coverage']['stop_reason'])}"])
    for gap in data["coverage"].get("gaps", []):
        lines.append(f"- {labels['gaps']}: {_text(gap)}")
    if data.get("validation_warnings"):
        lines.append("")
        lines.append("### Validation warnings")
        for warning in data["validation_warnings"]:
            lines.append(f"- {warning}")
    lines.append("")
    return "\n".join(lines)


def _source_links(source_ids: Iterable[str], source_by_id: dict[str, dict[str, Any]]) -> str:
    links: list[str] = []
    for source_id in source_ids:
        source = source_by_id[source_id]
        label = html.escape(source_id)
        links.append(f'<a href="#src-{html.escape(source_id)}">{label}</a>')
    return ", ".join(links) or "—"


def _html_claim_rows(claims: list[dict[str, Any]], source_by_id: dict[str, dict[str, Any]], include_rank: bool) -> str:
    rows: list[str] = []
    for claim in claims:
        searchable = " ".join([claim["method"], _text(claim.get("variant")), claim["score_display"], _text(claim.get("notes")), claim["claim_id"]]).casefold()
        cells: list[str] = []
        if include_rank:
            cells.append(f'<td data-sort="{claim.get("rank", 999999)}">{claim.get("rank", "—")}</td>')
        cells.extend([
            f"<td><strong>{html.escape(claim['method'])}</strong></td>",
            f"<td>{html.escape(_text(claim.get('variant'), '—'))}</td>",
            f'<td data-sort="{claim.get("score_numeric", "")}">{html.escape(claim["score_display"])}</td>',
            f"<td>{html.escape(_setting_text(claim.get('setting', {})) or '—')}</td>",
            f"<td>{_source_links(claim['evidence_source_ids'], source_by_id)}</td>",
            f"<td>{html.escape(_text(claim.get('notes'), '—'))}</td>",
            f"<td><code>{html.escape(claim['claim_id'])}</code></td>",
        ])
        rows.append(f'<tr data-search="{html.escape(searchable)}">' + "".join(cells) + "</tr>")
    return "".join(rows)


def _html_table(claims: list[dict[str, Any]], labels: dict[str, str], source_by_id: dict[str, dict[str, Any]], include_rank: bool) -> str:
    if not claims:
        return f"<p><em>{html.escape(labels['none'])}</em></p>"
    headers = ([labels["rank"]] if include_rank else []) + [labels["method"], labels["variant"], labels["score"], "Setting", labels["evidence"], labels["notes"], "claim_id"]
    head = "".join(f'<th tabindex="0">{html.escape(header)}</th>' for header in headers)
    return f'<div class="table-wrap"><table class="sortable"><thead><tr>{head}</tr></thead><tbody>{_html_claim_rows(claims, source_by_id, include_rank)}</tbody></table></div>'


def render_full_html(data: dict[str, Any]) -> str:
    run = data["research_run"]
    labels = _labels(run["language"])
    source_by_id = {source["source_id"]: source for source in data["sources"]}
    toc = "".join(f'<li><a href="#{html.escape(scope["scope_id"])}">{html.escape(scope["name"])}</a></li>' for scope in data["scopes"])
    scope_sections: list[str] = []
    for scope in data["scopes"]:
        axes = "".join(f"<li><code>{html.escape(_text(axis.get('name')))}</code>: {html.escape(_text(axis.get('value')))}</li>" for axis in scope.get("material_axes", [])) or f"<li>{html.escape(labels['none'])}</li>"
        main = _claims_for_scope(data, scope["scope_id"], "main")
        reference = _claims_for_scope(data, scope["scope_id"], "reference")
        unavailable = _claims_for_scope(data, scope["scope_id"], "unavailable")
        methods = "".join(f'<article class="method"><h4>{html.escape(claim["method"])}</h4><p>{html.escape(_text(claim.get("method_summary"), labels["none"]))}</p><code>{html.escape(claim["claim_id"])}</code></article>' for claim in main if claim.get("sota"))
        conflicts = [conflict for conflict in data["conflicts"] if conflict["scope_id"] == scope["scope_id"]]
        conflict_html = "".join(f'<li><strong>{html.escape(conflict["status"])}</strong>: {html.escape(conflict["summary"])}{(" — " + html.escape(_text(conflict.get("resolution")))) if _text(conflict.get("resolution")) else ""} <code>{html.escape(conflict["conflict_id"])}</code></li>' for conflict in conflicts) or f"<li>{html.escape(labels['none'])}</li>"
        scope_sections.append(f'''
<section id="{html.escape(scope['scope_id'])}">
  <h2>{html.escape(scope['name'])}</h2>
  <div class="scope-grid">
    <p><strong>{html.escape(labels['metric'])}:</strong> {html.escape(_metric_text(scope))}</p>
    <p><strong>Protocol:</strong> {html.escape(_text(scope.get('evaluation_protocol'), '—'))}</p>
  </div>
  <details><summary>Material comparison axes</summary><ul>{axes}</ul></details>
  <h3>{html.escape(labels['main'])}</h3>{_html_table(main, labels, source_by_id, True)}
  <details><summary>{html.escape(labels['reference'])} ({len(reference)})</summary>{_html_table(reference, labels, source_by_id, False)}</details>
  <details><summary>{html.escape(labels['unavailable'])} ({len(unavailable)})</summary>{_html_table(unavailable, labels, source_by_id, False)}</details>
  <details><summary>{html.escape(labels['methods'])}</summary><div class="method-grid">{methods or '<p>—</p>'}</div></details>
  <details><summary>{html.escape(labels['conflicts'])} ({len(conflicts)})</summary><ul>{conflict_html}</ul></details>
</section>''')
    sources_html = "".join(
        f'<article class="source" id="src-{html.escape(source["source_id"])}"><h3>{html.escape(source["title"])}</h3><p><code>{html.escape(source["source_id"])}</code> · {html.escape(source["authority"])} · {html.escape(_text(source.get("type")))}</p>'
        + (f'<p><a href="{html.escape(_text(source.get("url")), quote=True)}">{html.escape(_text(source.get("url")))}</a></p>' if source.get("url") else "")
        + f'<p>{html.escape(_text(source.get("notes"), ""))}</p></article>'
        for source in data["sources"]
    )
    gaps = "".join(f"<li>{html.escape(_text(gap))}</li>" for gap in data["coverage"].get("gaps", [])) or f"<li>{html.escape(labels['none'])}</li>"
    warnings = "".join(f"<li>{html.escape(warning)}</li>" for warning in data.get("validation_warnings", [])) or f"<li>{html.escape(labels['none'])}</li>"
    return f'''<!doctype html>
<html lang="{html.escape(run['language'])}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(labels['title'])}: {html.escape(run['topic'])}</title>
<style>
:root{{--ink:#172033;--muted:#667085;--line:#d8dee9;--panel:#f6f8fb;--accent:#2156a5;--main:#e9f6ee;--ref:#fff8e5}}
*{{box-sizing:border-box}} body{{margin:0;color:var(--ink);font:15px/1.55 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;background:white}}
header,main{{max-width:1240px;margin:auto;padding:24px}} header{{padding-top:42px;border-bottom:1px solid var(--line)}} h1{{font-size:clamp(28px,4vw,48px);line-height:1.1;margin:.2em 0}} h2{{margin-top:48px;border-bottom:2px solid var(--ink);padding-bottom:8px}} h3{{margin-top:28px}} a{{color:var(--accent)}} code{{font-size:.9em;background:#eef2f7;padding:2px 5px;border-radius:4px}}
.meta{{color:var(--muted)}} nav{{background:var(--panel);padding:16px 22px;border-radius:10px}} nav ul{{columns:2;gap:30px}} .filter{{position:sticky;top:0;z-index:5;background:white;padding:12px 0;border-bottom:1px solid var(--line)}} input{{width:min(520px,100%);padding:10px 12px;border:1px solid var(--line);border-radius:7px}}
.table-wrap{{overflow:auto;border:1px solid var(--line);border-radius:9px}} table{{width:100%;border-collapse:collapse;min-width:900px}} th,td{{padding:10px 12px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}} th{{position:sticky;top:0;background:#eef2f7;cursor:pointer;white-space:nowrap}} tbody tr:hover{{background:#f8fafc}}
details{{margin:16px 0;padding:10px 14px;border:1px solid var(--line);border-radius:9px}} summary{{cursor:pointer;font-weight:650}} .scope-grid,.method-grid{{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:12px}} .method,.source{{border:1px solid var(--line);border-radius:9px;padding:14px;background:white}} .source{{margin:10px 0}} footer{{margin-top:60px;padding:24px;color:var(--muted);border-top:1px solid var(--line)}}
@media(max-width:700px){{header,main{{padding:18px}} nav ul{{columns:1}}}}
</style>
</head>
<body>
<header><p class="meta">{html.escape(labels['cutoff'])}: {html.escape(run['evidence_checked_through'])}</p><h1>{html.escape(run['topic'])}</h1><p>{html.escape(run['query'])}</p></header>
<main>
<nav><strong>Contents</strong><ul>{toc}<li><a href="#coverage">{html.escape(labels['coverage'])}</a></li><li><a href="#sources">{html.escape(labels['sources'])}</a></li></ul></nav>
<div class="filter"><label>{html.escape(labels['filter'])}: <input id="filter" type="search" placeholder="method, score, claim ID..."></label></div>
{''.join(scope_sections)}
<section id="coverage"><h2>{html.escape(labels['coverage'])}</h2><p><strong>{html.escape(labels['stop'])}:</strong> {html.escape(_text(data['coverage']['stop_reason']))}</p><h3>{html.escape(labels['gaps'])}</h3><ul>{gaps}</ul><details><summary>Validation warnings</summary><ul>{warnings}</ul></details></section>
<section id="sources"><h2>{html.escape(labels['sources'])}</h2>{sources_html}</section>
<footer>run_id: <code>{html.escape(run['run_id'])}</code> · schema: <code>{SCHEMA_VERSION}</code> · generated: {html.escape(run['generated_at'])}</footer>
</main>
<script>
const filter=document.getElementById('filter');
filter.addEventListener('input',()=>{{const q=filter.value.toLowerCase();document.querySelectorAll('tbody tr').forEach(row=>row.hidden=!row.dataset.search.includes(q));}});
document.querySelectorAll('table.sortable th').forEach((th,index)=>{{th.addEventListener('click',()=>{{const table=th.closest('table');const body=table.tBodies[0];const rows=[...body.rows];const asc=th.dataset.asc!=='true';rows.sort((a,b)=>{{const av=a.cells[index]?.dataset.sort ?? a.cells[index]?.innerText ?? '';const bv=b.cells[index]?.dataset.sort ?? b.cells[index]?.innerText ?? '';const an=Number(av),bn=Number(bv);const cmp=Number.isNaN(an)||Number.isNaN(bn)?av.localeCompare(bv):an-bn;return asc?cmp:-cmp;}});rows.forEach(row=>body.appendChild(row));th.dataset.asc=String(asc);}});}});
</script>
</body></html>'''


def _sheet_title(name: str, used: set[str]) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "-", _text(name, "Scope"))[:31] or "Scope"
    candidate = clean
    suffix = 2
    while candidate in used:
        tail = f"-{suffix}"
        candidate = clean[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate)
    return candidate


def render_workbook(data: dict[str, Any], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ContractError("openpyxl is required for leaderboards.xlsx; install requirements.txt") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sota_fill = PatternFill("solid", fgColor="D9EAD3")
    reference_fill = PatternFill("solid", fgColor="FFF2CC")
    unavailable_fill = PatternFill("solid", fgColor="F4CCCC")
    used_titles: set[str] = set()
    source_by_id = {source["source_id"]: source for source in data["sources"]}

    def style_sheet(ws: Any, widths: dict[int, int]) -> None:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="top")
        for index, width in widths.items():
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    index = workbook.create_sheet("Index")
    used_titles.add("Index")
    index_headers = ["Scope", "Task", "Benchmark", "Metric", "Direction", "Main Results", "Reference Results", "Top Method", "Top Score", "Sheet", "Scope ID"]
    index.append(index_headers)

    for scope in data["scopes"]:
        title = _sheet_title(scope["name"], used_titles)
        claims = [claim for claim in data["claims"] if claim["scope_id"] == scope["scope_id"]]
        main = [claim for claim in claims if claim["section"] == "main"]
        reference = [claim for claim in claims if claim["section"] == "reference"]
        top = [claim for claim in main if claim.get("sota")]
        index.append([
            scope["name"], scope["task"], scope.get("benchmark"), _metric_text(scope), scope["metric"]["direction"], len(main), len(reference),
            "; ".join(claim["method"] for claim in top), "; ".join(claim["score_display"] for claim in top), title, scope["scope_id"],
        ])
        ws = workbook.create_sheet(title)
        headers = ["Rank", "Category", "Method", "Variant", "Score Numeric", "Score Display", "Paper / Result", "Year", "Setting", "Comparable", "SOTA", "Claim ID", "Evidence Source IDs", "Evidence URL", "Notes"]
        ws.append(headers)
        for claim in claims:
            first_url = ""
            for source_id in claim["evidence_source_ids"]:
                url = _text(source_by_id[source_id].get("url"))
                if url:
                    first_url = url
                    break
            ws.append([
                claim.get("rank"), claim["section"], claim["method"], claim.get("variant"), claim.get("score_numeric"), claim["score_display"],
                claim.get("paper_title"), claim.get("year"), _setting_text(claim.get("setting", {})), claim["comparability_complete"], claim["sota"],
                claim["claim_id"], ", ".join(claim["evidence_source_ids"]), first_url, claim.get("notes"),
            ])
            row = ws.max_row
            fill = sota_fill if claim.get("sota") else reference_fill if claim["section"] == "reference" else unavailable_fill if claim["section"] == "unavailable" else None
            if fill:
                for cell in ws[row]:
                    cell.fill = fill
            if first_url:
                ws.cell(row=row, column=14).hyperlink = first_url
                ws.cell(row=row, column=14).style = "Hyperlink"
        style_sheet(ws, {1: 9, 2: 13, 3: 24, 4: 20, 5: 14, 6: 14, 7: 34, 8: 9, 9: 32, 10: 12, 11: 9, 12: 32, 13: 34, 14: 42, 15: 42})

    style_sheet(index, {1: 30, 2: 25, 3: 24, 4: 28, 5: 12, 6: 14, 7: 17, 8: 28, 9: 18, 10: 28, 11: 34})
    workbook.save(path)


def _next_available_dir(base: Path) -> Path:
    if not base.exists() or not any((base / name).exists() for name in DEFAULT_ARTIFACTS):
        return base
    suffix = 2
    while True:
        candidate = base.with_name(f"{base.name}-{suffix}")
        if not candidate.exists() or not any((candidate / name).exists() for name in DEFAULT_ARTIFACTS):
            return candidate
        suffix += 1


def build_artifacts(data: dict[str, Any], output_root: Path | None = None, output_dir: Path | None = None) -> Path:
    """Write all artifacts through a staging directory and return the allocated run directory."""
    _require((output_root is None) != (output_dir is None), "Provide exactly one of output_root or output_dir")
    if output_dir is not None:
        base = output_dir
    else:
        base = output_root / data["research_run"]["topic_slug"]
    target = _next_available_dir(base)
    target.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".sota-finder-", dir=target.parent))
    try:
        (staging / "evidence.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (staging / "brief-report.md").write_text(render_brief_markdown(data), encoding="utf-8")
        (staging / "full-report.html").write_text(render_full_html(data), encoding="utf-8")
        render_workbook(data, staging / "leaderboards.xlsx")
        if target.exists():
            for artifact in DEFAULT_ARTIFACTS:
                _require(not (target / artifact).exists(), f"Refusing to overwrite existing artifact: {target / artifact}")
                shutil.move(str(staging / artifact), str(target / artifact))
            staging.rmdir()
        else:
            staging.rename(target)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return target


def load_input(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ContractError(f"Input file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ContractError(f"Invalid JSON in {path}: {exc}") from exc
    _require(isinstance(value, dict), "Research input must be a JSON object")
    return value


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build the standard SOTA Finder artifact bundle")
    parser.add_argument("--input", required=True, type=Path, help="Prepared research-input JSON")
    destination = parser.add_mutually_exclusive_group()
    destination.add_argument("--output-root", type=Path, default=Path("outputs"), help="Root for outputs/<topic-slug>/ (default: outputs)")
    destination.add_argument("--output-dir", type=Path, help="Explicit run directory; suffixes are still used to avoid overwrite")
    parser.add_argument("--validate-only", action="store_true", help="Normalize and validate without writing artifacts")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        normalized = normalize_bundle(load_input(args.input))
        if args.validate_only:
            print(json.dumps({"status": "valid", "schema_version": SCHEMA_VERSION, "warnings": normalized["validation_warnings"]}, ensure_ascii=False))
            return 0
        output_root = None if args.output_dir is not None else args.output_root
        target = build_artifacts(normalized, output_root=output_root, output_dir=args.output_dir)
        print(json.dumps({"status": "ok", "output_dir": str(target), "files": [str(target / name) for name in DEFAULT_ARTIFACTS], "warnings": normalized["validation_warnings"]}, ensure_ascii=False))
        return 0
    except ContractError as exc:
        print(json.dumps({"status": "error", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2
    except Exception as exc:
        print(json.dumps({"status": "error", "error": f"Unexpected failure: {exc}"}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
